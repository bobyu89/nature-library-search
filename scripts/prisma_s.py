"""PRISMA-S search appendix + checklist status from a structured search log (templates/searches.example.json).

usage: python -I prisma_s.py --log WORK/searches_sema.json --out WORK/PRISMA-S_sema.md [--docx WORK/PRISMA-S_sema.docx]

The Markdown has (1) a per-database table of every line with hits, limits, dates and platform, (2) other methods,
(3) a PRISMA-S 16-item checklist marked from what the log contains (✓ reported / ✗ missing / – not applicable), so
gaps are visible before submission. --docx needs python-docx (pip install python-docx); otherwise convert the .md.
"""
import argparse, json, sys

ITEMS = [
    (1, 'Database name', lambda L: all(d.get('name') for d in L['databases'])),
    (2, 'Multi-database searching (platform named)', lambda L: all(d.get('platform') for d in L['databases'])),
    (3, 'Study registries', lambda L: any('trial' in (o.get('name', '') + o.get('detail', '')).lower() or 'registry' in o.get('name', '').lower() or '試驗' in o.get('name', '') for o in L.get('other_methods', [])) or None),
    (4, 'Online resources and browsing', lambda L: any(o for o in L.get('other_methods', []) if 'scholar' in o.get('name', '').lower() or '網站' in o.get('name', '')) or None),
    (5, 'Citation searching', lambda L: any('引文' in o.get('name', '') or 'citation' in o.get('name', '').lower() for o in L.get('other_methods', [])) or None),
    (6, 'Contacts', lambda L: L.get('contacts') or None),
    (7, 'Other methods', lambda L: bool(L.get('other_methods'))),
    (8, 'Full search strategies (all lines, every database)', lambda L: all(d.get('lines') for d in L['databases'])),
    (9, 'Limits and restrictions', lambda L: all('limits' in d for d in L['databases'])),
    (10, 'Search filters (named, with source)', lambda L: all('filters' in d for d in L['databases'])),
    (11, 'Prior work (strategies adapted/reused)', lambda L: L.get('prior_work') or None),
    (12, 'Updates', lambda L: L.get('updates') or None),
    (13, 'Dates of searches', lambda L: all(d.get('searched') for d in L['databases'])),
    (14, 'Peer review of strategy', lambda L: bool(L.get('peer_review_of_strategy'))),
    (15, 'Total records per database', lambda L: all(d.get('records_exported') is not None for d in L['databases'])),
    (16, 'Deduplication method', lambda L: bool((L.get('dedup') or {}).get('method'))),
]


def mark(v):
    return '✓' if v is True or (v and v is not None) else ('–' if v is None else '✗')


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--log', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--docx')
    a = ap.parse_args()
    L = json.load(open(a.log, encoding='utf8'))
    md = ['# 補充資料：完整檢索策略（PRISMA-S）', '', '臨床問題：%s' % L.get('question', ''), '',
          '出版期間：%s；語言限制：%s；檢索者：%s；檢索策略審閱：%s' % (L.get('date_range', ''), L.get('language_limit', ''), L.get('searcher', ''), L.get('peer_review_of_strategy', '')), '']
    for d in L['databases']:
        md += ['## %s（%s；%s）' % (d['name'], d.get('platform', ''), d.get('access', '')),
               '檢索日期：%s；書目匯出：%s；限制：%s；過濾：%s；匯出筆數：%s' % (d.get('searched', ''), d.get('exported', ''), d.get('limits', ''), d.get('filters', ''), d.get('records_exported', '')), '',
               '| # | 檢索式 | 筆數 |', '|---|---|---|']
        md += ['| %s | %s | %s |' % (x['n'], x['query'].replace('|', '\\|'), '{:,}'.format(x['hits']) if isinstance(x['hits'], int) else x['hits']) for x in d.get('lines', [])]
        if d.get('notes'):
            md += ['', '註：' + d['notes']]
        md.append('')
    if L.get('other_methods'):
        md += ['## 其他方法', '| 方法 | 日期 | 內容 | 筆數 | 新增納入 |', '|---|---|---|---|---|']
        md += ['| %s | %s | %s | %s | %s |' % (o.get('name', ''), o.get('date', ''), o.get('detail', ''), o.get('hits', ''), o.get('added', '')) for o in L['other_methods']]
        md.append('')
    md += ['## 去除重複', '%s；移除 %s 筆。' % (L.get('dedup', {}).get('method', ''), L.get('dedup', {}).get('removed', '')), '']
    if L.get('seed_check'):
        md += ['## 檢索策略驗證', L['seed_check'], '']
    md += ['## PRISMA-S 檢核（✓ 已報告；✗ 缺；– 不適用或未執行）', '| # | 項目 | 狀態 |', '|---|---|---|']
    md += ['| %d | %s | %s |' % (n, name, mark(f(L))) for n, name, f in ITEMS]
    open(a.out, 'w', encoding='utf8').write('\n'.join(md) + '\n')
    print('written', a.out)
    missing = [name for n, name, f in ITEMS if mark(f(L)) == '✗']
    if missing:
        print('missing PRISMA-S items:', '; '.join(missing))
    if a.docx:
        try:
            import docx
        except ImportError:
            print('python-docx not installed; skip --docx')
            return
        doc = docx.Document()
        rows = []

        def flush():
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = 'Table Grid'
                for i, r in enumerate(rows):
                    for j, c in enumerate(r[:len(rows[0])]):
                        t.cell(i, j).text = c.replace('\\|', '|')
                rows.clear()

        for line in md:
            if line.startswith('|'):
                if not line.startswith('|---'):
                    rows.append([c.strip() for c in line.strip().strip('|').split(' | ')])
                continue
            flush()
            if line.startswith('# '):
                doc.add_heading(line[2:], 1)
            elif line.startswith('## '):
                doc.add_heading(line[3:], 2)
            elif line:
                doc.add_paragraph(line)
        flush()
        doc.save(a.docx)
        print('written', a.docx)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
